from docx import Document
from pathlib import Path

models = Path('Models')
for p in sorted(models.glob('*.docx')):
    try:
        doc = Document(str(p))
        text = '\n'.join([para.text for para in doc.paragraphs])
        # crude find of {{ ... }} occurrences
        vars_found = set()
        import re
        for m in re.finditer(r"\{\{\s*([^\}]+)\s*\}\}", text):
            vars_found.add(m.group(1).strip())
        print(p.name)
        if vars_found:
            for v in sorted(vars_found):
                print('  ', v)
        else:
            print('   (no {{ }} placeholders found in paragraph text)')
    except Exception as e:
        print('Failed to read', p.name, e)
